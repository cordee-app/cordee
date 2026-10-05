"""RAG integration helpers for AIngel — provider-agnostic legal retrieval.

The RAG factory lives OUTSIDE ``AINGEL_PROJECTS_ROOT`` at ``/opt/RAG`` (Qdrant
:6333, HTTP API :8010) and is shared across projects (1:N). This module is the
vault-side brain that decides *when* to retrieve (legal intent detection),
*what* to inline into the execution prompt (pre-fetch), and *how* to label the
result so the UI can show whether a run grounded its answer in the library or
in the model's generic knowledge.

Three responsibilities (mirrors the file-auto-catch pattern in
``prompt_builder.py`` but for the shared legal library rather than project
Working Docs):

1. ``detect_rag_intent`` — should this task consult the RAG library at all?
2. ``prefetch_citations``  — run the retrieval and format an inline context block.
3. ``provenance_label``    — classify a run for the UI badge (RAG-PREFETCHED /
   RAG-TOOL / RAG-REQUIRED / RAG-OFF).

Retrieval quality (hybrid BGE-M3 + BM25 + article boost + RRF, optional
cross-encoder rerank, optional graph expansion) lives in the factory
(``/opt/RAG/scripts/rag_search.py``) — the vault only passes flags through.
"""
import json
import logging
import os
import re

_log = logging.getLogger(__name__)

# Known factory corpora (authoritative list is /opt/RAG/corpora/_registry.yaml).
# `polish_general` is structurally ready but NOT indexed yet — keep it out of
# automatic detection so we never offer an empty library.
RAG_CORPORA = ('railway', 'polish_general_law')

# Human-facing labels for the RAG corpus dropdowns (kept in sync with
# /opt/RAG/corpora/_registry.yaml names). Used by the UI via /api/config.
RAG_CORPORA_LABELS = {
    'railway': 'railway — EU→PL railway law',
    'polish_general_law': 'polish_general_law — Polish commercial law (KSH)',
}

# ── Block D1: rag-index tag — host-local opt-in (Phase 4) ───────────────────
# Files tagged ``rag-index`` are indexed locally under /opt/RAG (host-local,
# no SCW egress). Chunk size mirrors batch mode (~280k chars). Retrieval is
# ONLY injected when the task opts in via context_refs OR "rag:" keyword.

RAG_INDEX_TAG = 'rag-index'
_RAG_CHUNK_CHARS = 280_000
_RAG_LOCAL_ROOT = '/opt/RAG'
_RAG_LOCAL_INDICES_DIR = os.path.join(_RAG_LOCAL_ROOT, 'local_indices')

# 41 DU hold-out PDFs — QA validation set, deliberately NOT indexed.
# Source: a client project's Working Documents (ls DU*.pdf = 41)
# + 1 publ_ variant (42 total PDFs with du in name; 41 start with DU).
_HOLDOUT_BASENAMES_LOWER = frozenset(s.lower() for s in (
    'DU2002521000-sig.pdf',
    'DU20202800-sig.pdf',
    'DU20210600-sig.pdf',
    'DU20210700-sig.pdf',
    'DU20220100-sig.pdf',
    'DU20220600_sig.pdf',
    'DU20221400-sig.pdf',
    'DU20221600-sig.pdf',
    'DU20221700-sig.pdf',
    'DU20221800-sig.pdf',
    'DU20222100-sig.pdf',
    'DU202237000_sig.pdf',
    'DU202305000-sig.pdf',
    'DU202323000-sig.pdf',
    'DU202405000-sig.pdf',
    'DU202420000-signed.pdf',
    'DU202422000-sig.pdf',
    'DU202423000-sig.pdf',
    'DU202424000-sig.pdf',
    'DU20242500-sig.pdf',
    'DU202426000-sig.pdf',
    'DU202427000-sig.pdf',
    'DU202433000-sig.pdf',
    'DU202434000-sig.pdf',
    'DU202520001-sig.pdf',
    'DU202522000-sig.pdf',
    'DU202523000-sig.pdf',
    'DU202524000-sig.pdf',
    'DU202526000-sig.pdf',
    'DU202528000-sig.pdf',
    'DU202530000-sig.pdf',
    'DU202531000-sig.pdf',
    'DU_20211000-sig.pdf',
    'DU_20211100-sig.pdf',
    'DU_20211200-sig.pdf',
    'DU_202322000-sig.pdf',
    'DU_20240300-sig.pdf',
    'DU_202413000-sig.pdf',
    'DU_202414000-sig.pdf',
    'DU_2025004000-sig.pdf',
    'DU_202518000-sig.pdf',
    # plus the publ_ variant (same DU set, different prefix)
    'publ_DU202421000_podpis.pdf',
))

# Which corpus a query should target, keyed on distinctive terms.
_CORPUS_KEYWORDS = {
    'railway': (
        'kolej', 'railway', 'rail', 'trains', 'pociąg', 'maszynista',
        'przewoźnik', 'infrastruktura kolejowa', 'plk', 'utk', 'era',
        'świadectwo maszynisty', 'licencja przewoźnika', 'tabor', 'wagony',
        'szyny', 'tory', 'interoperacyjność', 'safety', 'transport kolejowy',
    ),
    'polish_general_law': (
        'spółka', 'spolka', 'spółki', 'spolki', 'akcyjna', 'akcjonariusz',
        'udziałowiec', 'udział', 'zarząd', 'kodeks spółek', 'prawo spółek',
        'spółka z o.o.', 'spółka akcyjna', 'prosta spółka', 'komandytowa',
        'kapitał zakładowy', 'zgromadzenie wspólników', 'walne zgromadzenie',
        'firma', 'rejestr przedsiębiorców', 'krs', 'spółka jawna', 'spółka partnerska',
        # VAT (Ustawa o podatku od towarów i usług, Dz.U. 2025 poz.775)
        'vat', 'podatek od towarów', 'podatku od towarów', 'towarów i usług',
        'faktura', 'faktury', 'podatnik', 'podatnika', 'podatnicy', 'vat-ue',
        'stawka podatku', 'opodatkowanie', 'zwolnienie z vat', 'odliczenie vat',
        'deklaracja vat', 'nowe środki transportu', 'złoto inwestycyjne',
    ),
}

def list_corpora():
    """Return [{id, label}] for the corpora that are actually usable."""
    return [
        {'id': cid, 'label': RAG_CORPORA_LABELS.get(cid, cid)}
        for cid in RAG_CORPORA
    ]

# Regex markers that a legal task (Polish/EU statutes/regulations) is being asked.
# ELI/CELEX are bare acronyms and MUST be word-boundary-anchored — unanchored,
# ELI matches "eli" inside ordinary words like "guidelines" or "reliable"
# (regression: task #10001145 — an example phrase "EU tendering guidelines" in
# the task's own instructions false-triggered RAG via "guidELInes").
_LEGAL_REGEX = re.compile(
    r'Dz\.U\.|\bELI\b|\bCELEX\b|eli\.gov\.pl|\bArt\.?\s*\d+[a-z]*|art\.?\s*\d+[a-z]*'
    r'|\bust\.|\bustaw|rozporządzeni|rozporzadzeni|dyrekty|dyrektyw'
    r'|\bustawa\b|\bprawo\b|kodeks|zapisy prawne|przepis'
    r'|licencj|świadectwo maszynisty|swiadectwo maszynisty|koncesj|regulamin',
    re.IGNORECASE,
)

# How many chunks to retrieve by default; capped small so the inline block stays
# inside the prompt budget (see _RAG_INLINE_CHAR_CAP).
_RAG_TOP_K = 6

# Inline block cap — a single execution prompt must stay well under the ~40 kB
# Scaleway body limit and within the model context window.
_RAG_INLINE_CHAR_CAP = 4000

# Max length of the focused RAG query. A long multi-step task description dilutes
# semantic retrieval (BM25/dense see a wall of instructions instead of the legal
# provisions the task needs), so we extract a compact legal question instead.
_RAG_QUERY_CHAR_CAP = 300

# Matches a specific article reference, e.g. "art. 202 § 6", "Art. 116", "art. 22b ust. 4".
_ARTICLE_REF = re.compile(
    r'\b(?:art\.?|artyku[łl])\s*\d+[a-z]*(?:\s*[§\u00a7]\s*\d+[a-z]*)?'
    r'(?:\s*(?:ust\.?|ustęp|ustep)\s*\d+[a-z]*)?(?:\s*(?:pkt|punkt)\s*\d+[a-z]*)?',
    re.IGNORECASE,
)

# Statute / act names worth keeping in a focused query.
_STATUTE_TERMS = (
    'kodeks spółek handlowych', 'kodeks spółek', 'k.s.h.', 'ksh',
    'kodeks cywilny', 'k.c.', 'ordynacja podatkowa', 'o.p.', 'ustawa o podatku',
    'ustawa o rachunkowości', 'prawo upadłościowe', 'prawo restrukturyzacyjne',
    'prawo kolejowe', 'ustawa o transporcie kolejowym', 'rozporządzenie', 'dyrektywa',
)


def _extract_legal_query(desc, title):
    """Build a focused RAG query from a (possibly long) task description.

    Sending the whole multi-step instruction block as the retrieval query dilutes
    semantic search — the model pulls generic articles instead of the specific
    provisions the task needs (observed: task 10001105 needed art. 202 § 6 k.s.h.
    but retrieved art. 402/580/550). Instead we extract the legal substance:

      1. Every explicit article reference (art. 202 § 6, Art. 116, ...).
      2. Statute/act names (kodeks spółek handlowych, ordynacja podatkowa, ...).
      3. The most legally-relevant sentences (those matching legal markers).

    Returns a compact query string ('' if nothing legal is found).
    """
    blob = f'{desc or ""} {title or ""}'
    if not blob.strip():
        return ''

    # Strip placeholder / illustrative citations and OCR-flag examples BEFORE
    # extracting article references. Task descriptions often contain *examples*
    # of how to cite or flag OCR errors — e.g. "(Dz.U. 2023 poz. 1234, ELI:
    # [link])" or "[OCR: możliwy błąd w tekście: 'art. 12' → weryfikować z
    # oryginałem]". These are NOT real legal references the task needs, but the
    # article regex would otherwise pick up the fake "art. 12" and pollute the
    # prefetch query (observed: task 10001109 retrieved art. 12 OP instead of
    # art. 116 § 2 OP). Remove them so only genuine provisions survive.
    blob = re.sub(
        r'\(Dz\.U\.\s*\d{4}\s+poz\.\s*\d+[^)]*\)', ' ', blob, flags=re.IGNORECASE)
    blob = re.sub(
        r'\[OCR:[^\]]*\]', ' ', blob, flags=re.IGNORECASE)
    blob = re.sub(r'\s+', ' ', blob)

    # 1. Article references — the strongest signal. Keep them verbatim.
    articles = _ARTICLE_REF.findall(blob)
    # 2. Statute names.
    low = blob.lower()
    statutes = [t for t in _STATUTE_TERMS if t in low]

    # 3. Legal-relevant sentences (fallback / enrichment).
    sentences = re.split(r'(?<=[.!?])\s+', blob)
    legal_sents = [s.strip() for s in sentences if _LEGAL_REGEX.search(s)]

    parts = []
    if articles:
        parts.append(' '.join(dict.fromkeys(a.strip() for a in articles)))
    if statutes:
        parts.append(' '.join(dict.fromkeys(statutes)))
    if legal_sents:
        # Prefer the shortest legal sentences (most likely to be the actual question).
        legal_sents.sort(key=len)
        parts.append(' '.join(legal_sents[:3]))

    query = ' '.join(parts).strip()
    if not query:
        # Nothing legal extracted — fall back to the title (short, focused).
        query = (title or desc or '').strip()
    return query[:_RAG_QUERY_CHAR_CAP]


def _project_enables_rag(project):
    """RAG is available only when the project opts in (per-project library)."""
    try:
        return bool((project or {}).get('use_rag'))
    except Exception:
        return False


def _task_requires_rag(task):
    """Per-task explicit opt-in (the user marks a task as needing the library)."""
    try:
        return bool((task or {}).get('requires_rag'))
    except Exception:
        return False


def _role_is_legal(task):
    try:
        role = (task or {}).get('role_name') or ''
    except Exception:
        role = ''
    return 'legal' in (role or '').lower()


def detect_rag_intent(task, project):
    """Return a retrieval intent dict if this task should consult RAG, else None.

    Intent fires when:
      * the project has RAG enabled (``use_rag``) AND
      * the task is explicitly flagged ``requires_rag`` OR auto-intent matches
        (Legal project_type, legal role, or legal markers in the description).

    Returns ``{corpus_id, query, top_k, rerank, graph, reason, requires_rag}``
    or ``None`` when RAG should not be consulted.
    """
    if not _project_enables_rag(project):
        return None

    task = task or {}
    desc = (task.get('description') or '').strip()
    title = (task.get('title') or '').strip()
    project_type = ((project or {}).get('project_type') or '').lower()

    auto = bool(
        project_type == 'legal'
        or _role_is_legal(task)
        or bool(_LEGAL_REGEX.search(desc or title))
    )
    requires = _task_requires_rag(task)
    if not (requires or auto):
        return None

    # Corpus selection: an explicit task corpus override wins; else railway terms
    # → railway; otherwise the first populated corpus as the legal fallback.
    blob = (desc + ' ' + title).lower()
    explicit_corpus = ((task or {}).get('corpus_id') or '').strip()
    if explicit_corpus in RAG_CORPORA:
        corpus_id = explicit_corpus
    else:
        corpus_id = next(
            (cid for cid, kws in _CORPUS_KEYWORDS.items() if any(k in blob for k in kws)),
            RAG_CORPORA[0],
        )

    # Focused legal query — NOT the whole task description (which dilutes
    # retrieval). Extract the specific articles/statutes/legal sentences.
    query = _extract_legal_query(desc, title)
    if not query:
        return None

    return {
        'corpus_id': corpus_id,
        'query': query,
        'top_k': _RAG_TOP_K,
        # rerank (cross-encoder) is factory-side + CPU-slow — off by default.
        'rerank': False,
        # graph expansion (L1→L2→L3 hierarchy) is cheap and improves the
        # cross-layer picture for legal reasoning — on for legal intent.
        'graph': True,
        'reason': 'explicit requires_rag' if requires else 'auto-detected legal intent',
        'requires_rag': requires,
    }


def _safe(v):
    if v is None:
        return None
    if isinstance(v, float):
        return v if (v == v and v != float('inf') and v != float('-inf')) else None
    return v


def _format_citation(c):
    """Render one citation chunk into a compact legal citation line."""
    layer = (c.get('layer') or '').strip()
    src = (c.get('dz_u') or c.get('celex') or '').strip()
    eli = (c.get('eli') or '').strip()
    article = (c.get('article') or '').strip()
    ustep = (c.get('ustep') or '').strip()
    punkt = (c.get('punkt') or '').strip()

    ref = src
    parts = []
    if article:
        parts.append(f'Art.{article}')
    if ustep:
        parts.append(f'ust.{ustep}')
    if punkt:
        parts.append(f'pkt {punkt}')
    loc = '. '.join(parts)
    if loc and ref:
        ref = f'{ref} {loc}'
    elif loc:
        ref = loc

    label = f'[{layer}] {ref}'.strip()
    if eli:
        label += f' — {eli}'
    extract = (c.get('extract') or '').strip()
    if extract:
        label += f'\n{extract}'
    return label


def _format_block(intent, resp):
    """Build the inline 'Retrieved legal context' markdown block for the prompt."""
    citations = (resp or {}).get('citations') or []
    corpus = intent.get('corpus_id', '')
    via = (resp or {}).get('via') or ''
    lines = [f'## Retrieved legal context (RAG: {corpus}, {len(citations)} chunk(s))']
    lines.append(
        'This is the authoritative legal library (factory /opt/RAG). Answer statutes/'
        'regulations ONLY from these excerpts. Cite Dz.U./CELEX + ELI + the verbatim '
        'extract. If the answer is not here, say "Brak w dostarczonych aktach" — do '
        'not answer from memory or invent citations.'
    )
    lines.append(
        'IMPORTANT: before you cite ANY article (e.g. art. 116 § 2 Ordynacji '
        'podatkowej, art. 202 § 6 k.s.h.), call the rag_query tool with that exact '
        'article reference to fetch its verbatim text. The excerpts above are only a '
        'pre-fetch and may not include the specific article you need. Do NOT write '
        '"Brak w dostarczonych aktach" for an article until you have actually queried '
        'it via rag_query and it returned nothing.'
    )
    for c in citations:
        lines.append(_format_citation(c))

    # Graph trace (L1→L2→L3 hierarchy relations) when the factory returned one.
    retrieval = (resp or {}).get('retrieval') or {}
    graph_trace = retrieval.get('graph_trace') or []
    if graph_trace:
        lines.append('POWIĄZANIA (hierarchia L1>L2>L3):')
        for t in graph_trace:
            lines.append(f'  - {t}')
    block = '\n\n'.join(lines)
    return block[:_RAG_INLINE_CHAR_CAP]


def prefetch_citations(intent):
    """Run retrieval for an intent and return ``(block_text, provenance_dict)``.

    ``block_text`` is ready to inline into the prompt ('' on failure). The
    ``provenance_dict`` records *how* the retrieval happened so the UI can show
    a transparent RAG badge.
    """
    if not intent:
        return '', {}
    try:
        from agent_tools import rag_query
        resp = rag_query(
            corpus_id=intent.get('corpus_id', ''),
            query=intent.get('query', ''),
            top_k=intent.get('top_k', _RAG_TOP_K),
            filters={},
            rerank=bool(intent.get('rerank', False)),
            graph=bool(intent.get('graph', False)),
        )
    except Exception:
        return '', {}

    citations = (resp or {}).get('citations') or []
    provenance = {
        'via': _safe((resp or {}).get('via')),
        'corpus_id': intent.get('corpus_id'),
        'collection': _safe((resp or {}).get('collection')),
        'n_hits': len(citations),
        'query': intent.get('query'),
        'rerank': bool(intent.get('rerank')),
        'graph': bool(intent.get('graph')),
        'retrieval': _safe((resp or {}).get('retrieval')),
        'citation_ids': [_safe(c.get('chunk_id')) for c in citations],
        'dz_u_refs': [c.get('dz_u') for c in citations if c.get('dz_u')],
        'celexes': [c.get('celex') for c in citations if c.get('celex')],
    }
    if (resp or {}).get('error'):
        provenance['error'] = resp.get('error')
    if not citations:
        return '', provenance

    block = _format_block(intent, resp)
    return block, provenance


def provenance_label(task, project, prefetch_ok, tool_calls=None):
    """Classify a run into one of the UI badge labels.

    Args:
        prefetch_ok: bool — the pre-fetch inlined non-empty citations.
        tool_calls: optional iterable of tool names the model invoked mid-run.

    Returns one of: RAG-PREFETCHED | RAG-TOOL | RAG-REQUIRED | RAG-OFF.
    """
    tool_calls = tool_calls or []
    called_rag = any('rag_query' == (t or '') for t in tool_calls)
    requires = _task_requires_rag(task)
    enabled = _project_enables_rag(project)

    if called_rag:
        return 'RAG-TOOL'
    if prefetch_ok:
        return 'RAG-PREFETCHED'
    if requires:
        return 'RAG-REQUIRED'   # expected but missed (empty / library down)
    if enabled:
        return 'RAG-OFF'        # library available but no intent for this task
    return 'RAG-OFF'


# ── Block D1: rag-index host-local helpers ──────────────────────────────────

def is_holdout_file(rel):
    """Return True if *rel* is in the 41 DU hold-out set (case-insensitive).

    Comparison is on basename or full rel lowercased, plus a pattern fallback
    (any basename starting with ``du`` and ending ``.pdf``) so a renamed copy
    is also protected.
    """
    if not rel:
        return False
    base = os.path.basename(rel.strip()).lower()
    low_rel = rel.strip().lower()
    if base in _HOLDOUT_BASENAMES_LOWER or low_rel in _HOLDOUT_BASENAMES_LOWER:
        return True
    # Pattern fallback — any DU PDF, including publ_DU variant
    if base.startswith('du') and base.endswith('.pdf'):
        return True
    if 'du20' in base and base.endswith('.pdf'):
        return True
    return False


def _rag_extract_text(full_path, max_chars=500_000):
    """Extract text from *full_path* (host-local, no SCW). Returns str or ''."""
    if not full_path or not os.path.isfile(full_path):
        return ''
    ext = os.path.splitext(full_path)[1].lower()
    try:
        if ext == '.pdf':
            try:
                import pdfplumber
                parts = []
                with pdfplumber.open(full_path) as pdf:
                    for page in pdf.pages:
                        t = page.extract_text() or ''
                        if t.strip():
                            parts.append(t)
                txt = '\n'.join(parts)
                if txt.strip():
                    return txt[:max_chars]
            except Exception:
                pass
        elif ext == '.docx':
            try:
                import docx
                doc = docx.Document(full_path)
                parts = [p.text.strip() for p in doc.paragraphs if p.text.strip()]
                for table in doc.tables:
                    for row in table.rows:
                        cells = [c.text.strip() for c in row.cells if c.text.strip()]
                        if cells:
                            parts.append(' | '.join(cells))
                txt = '\n'.join(parts)
                if txt.strip():
                    return txt[:max_chars]
            except Exception:
                pass
        elif ext == '.xlsx':
            try:
                import openpyxl
                wb = openpyxl.load_workbook(full_path, read_only=True, data_only=True)
                parts = []
                for ws in wb.worksheets:
                    for row in ws.iter_rows(values_only=True):
                        cells = [str(c) for c in row if c is not None]
                        if cells:
                            parts.append(' | '.join(cells))
                wb.close()
                txt = '\n'.join(parts)
                if txt.strip():
                    return txt[:max_chars]
            except Exception:
                pass
        elif ext == '.pptx':
            try:
                from pptx import Presentation
                prs = Presentation(full_path)
                parts = []
                for i, slide in enumerate(prs.slides, 1):
                    slide_texts = []
                    for shape in slide.shapes:
                        if shape.has_text_frame:
                            for para in shape.text_frame.paragraphs:
                                t = para.text.strip()
                                if t:
                                    slide_texts.append(t)
                        if shape.has_table:
                            for row in shape.table.rows:
                                cells = [cell.text.strip() for cell in row.cells if cell.text.strip()]
                                if cells:
                                    slide_texts.append(' | '.join(cells))
                    if slide_texts:
                        parts.append(f'--- Slide {i} ---\n' + '\n'.join(slide_texts))
                txt = '\n'.join(parts)
                if txt.strip():
                    return txt[:max_chars]
            except Exception:
                pass
        # Fallback: utf-8 text read
        with open(full_path, 'r', encoding='utf-8', errors='replace') as fh:
            txt = fh.read()
            return txt[:max_chars] if txt.strip() else ''
    except Exception:
        return ''
    return ''


def _rag_split_text(text, chunk_chars=_RAG_CHUNK_CHARS):
    """Split *text* into <=chunk_chars pieces, preferring newline boundaries."""
    text = text or ''
    if len(text) <= chunk_chars:
        return [text] if text.strip() else []
    chunks, start, n = [], 0, len(text)
    while start < n:
        end = min(start + chunk_chars, n)
        if end < n:
            nl = text.rfind('\n', start, end)
            if nl > start:
                end = nl + 1
        chunks.append(text[start:end])
        start = end
    return chunks


def _rag_local_index_path(project_path, rel):
    """Return host-local index file for *rel* under /opt/RAG/local_indices/<slug>/."""
    slug = os.path.basename((project_path or '').rstrip('/')).replace(' ', '_') or 'default'
    # Sanitize rel for filesystem — keep subdirs but strip traversal
    safe_rel = rel.strip().replace(os.sep, '/').lstrip('/')
    safe_rel = safe_rel.replace('..', '_')
    return os.path.join(_RAG_LOCAL_INDICES_DIR, slug, safe_rel + '.chunks.json')


def reindex_file(project_path, rel):
    """Re-embed a single ``rag-index`` file locally (host-local, no SCW).

    * Respects the 41 DU hold-out — skipped with log ``hold-out skipped``.
    * Chunks at ~280k chars (mirrors batch mode).
    * EU-only: host-local embedding, no US model is called (BGE-M3 local if
      available, otherwise plain chunk storage — still local).
    * Best-effort: never raises to caller.
    Returns True on success, False on skip/failure.
    """
    rel_n = (rel or '').strip().replace(os.sep, '/')
    if not rel_n or not project_path:
        return False
    if is_holdout_file(rel_n):
        _log.info('hold-out skipped: %s', rel_n)
        # Also print for manual test visibility (journalctl)
        print(f'[rag-index] hold-out skipped: {rel_n}')
        return False
    # EU-only gate: local RAG is allowed (host-local), but log that we are
    # staying on-host and not calling a US embedding model.
    try:
        import agent_config as _cfg
        proj = None
        # Try to resolve project row for eu_only flag
        try:
            import agent_db as _db
            for p in _db.get_projects():
                if p.get('path') == project_path:
                    proj = p
                    break
        except Exception:
            proj = None
        if _cfg.eu_only_for(proj):
            _log.info('[rag-index] eu_only project — using host-local /opt/RAG only (no US model)')
    except Exception:
        pass

    full = os.path.join(project_path, rel_n)
    if not os.path.isfile(full):
        # Also try basename search under writable variants (same as prompt_builder)
        found = None
        if '/' not in rel_n:
            for folder in ('Working Documents', 'Working Docs', 'My Docs', 'working-docs', 'docs'):
                cand = os.path.join(project_path, folder, rel_n)
                if os.path.isfile(cand):
                    full = cand
                    found = True
                    break
        if not found and not os.path.isfile(full):
            _log.warning('[rag-index] file not found for reindex: %s', rel_n)
            return False

    text = _rag_extract_text(full)
    if not text or not text.strip():
        _log.warning('[rag-index] empty text, skip: %s', rel_n)
        return False

    chunks = _rag_split_text(text, _RAG_CHUNK_CHARS)
    if not chunks:
        _log.warning('[rag-index] no chunks: %s', rel_n)
        return False

    # Persist chunks locally (host-local /opt/RAG, no SCW)
    idx_path = _rag_local_index_path(project_path, rel_n)
    try:
        os.makedirs(os.path.dirname(idx_path), exist_ok=True)
        payload = {
            'rel': rel_n,
            'project_path': project_path,
            'chunks': [{'chunk_id': f'{rel_n}::chunk-{i}', 'text': ch} for i, ch in enumerate(chunks, 1)],
            'n_chunks': len(chunks),
            'chars': len(text),
        }
        with open(idx_path, 'w', encoding='utf-8') as fh:
            json.dump(payload, fh, ensure_ascii=False, indent=2)
        _log.info('[rag-index] reindexed %s → %s (%d chunks, %d chars) host-local /opt/RAG', rel_n, idx_path, len(chunks), len(text))
        print(f'[rag-index] reindexed {rel_n} ({len(chunks)} chunks, {len(text)} chars) → {idx_path} (host-local)')
    except Exception as e:
        _log.warning('[rag-index] local persist failed for %s: %s', rel_n, e)
        return False

    # Best-effort Qdrant upsert (local Qdrant :6333, same host, no SCW). If
    # Qdrant is down, the JSON above is already the retrieval source.
    try:
        from qdrant_client import QdrantClient  # type: ignore
        from qdrant_client.models import Distance, VectorParams, PointStruct  # type: ignore
        import uuid as _uuid
        # Try local embedding via sentence-transformers (BGE-M3, EU model on-host)
        try:
            from sentence_transformers import SentenceTransformer  # type: ignore
            _model = SentenceTransformer('BAAI/bge-m3', trust_remote_code=True, device='cpu')
            vecs = _model.encode([c[:1200] for c in chunks], normalize_embeddings=True, batch_size=8).tolist()
        except Exception as _emb_e:
            _log.info('[rag-index] local embed skipped (no BGE-M3): %s', _emb_e)
            vecs = None
        if vecs:
            _slug = os.path.basename((project_path or '').rstrip('/')).replace(' ', '_') or 'default'
            _coll = f'rag_index_{_slug}'[:63]
            _client = QdrantClient(host='localhost', port=6333, timeout=10)
            cols = [c.name for c in _client.get_collections().collections]
            if _coll not in cols:
                _client.create_collection(collection_name=_coll, vectors_config=VectorParams(size=len(vecs[0]), distance=Distance.COSINE))
            _ns = _uuid.NAMESPACE_DNS
            points = []
            for ch, vec in zip(chunks, vecs):
                # chunk_id already scoped by rel; make point id stable
                pid = str(_uuid.uuid5(_ns, f'{rel_n}::{ch[:64]}'))
                points.append(PointStruct(id=pid, vector=vec, payload={'rel': rel_n, 'text': ch[:1200], 'chunk_id': f'{rel_n}::chunk'}))
            # Upsert in batches
            for i in range(0, len(points), 50):
                _client.upsert(collection_name=_coll, points=points[i:i+50])
            _log.info('[rag-index] qdrant upsert %s (%d points) host-local', _coll, len(points))
    except Exception as _q_e:
        _log.info('[rag-index] qdrant local upsert skipped: %s', _q_e)

    return True


def _has_rag_keyword(task):
    """Return True if description contains ``rag:`` (case-insensitive)."""
    desc = (task or {}).get('description') or ''
    return 'rag:' in desc.lower()


def _task_rag_index_refs(task, project_path):
    """Return list of context_refs that have ``rag-index`` tag (host-local check)."""
    raw = (task or {}).get('context_refs')
    refs = []
    if isinstance(raw, str):
        try:
            parsed = json.loads(raw) if raw.strip() else []
            if isinstance(parsed, list):
                refs = parsed
        except Exception:
            refs = []
    elif isinstance(raw, list):
        refs = raw
    norm = []
    for r in (refs or []):
        if not isinstance(r, str):
            continue
        s = r.strip().replace(os.sep, '/')
        if s and not os.path.isabs(s) and '..' not in s.split('/'):
            if s not in norm:
                norm.append(s)
    if not norm or not project_path:
        return []
    try:
        import agent_db as _db
        tags_map = _db.get_file_tags(project_path) or {}
    except Exception:
        tags_map = {}
    out = []
    for rel in norm:
        meta = tags_map.get(rel)
        if not meta:
            # Also try basename match (file may be referenced without Working Documents/ prefix)
            base = os.path.basename(rel)
            for k, v in tags_map.items():
                if os.path.basename(k) == base and any(str(t).lower() == RAG_INDEX_TAG for t in (v.get('tags') or [])):
                    out.append(rel)
                    break
            continue
        tags = [str(t).lower() for t in (meta.get('tags') or [])]
        if RAG_INDEX_TAG in tags:
            out.append(rel)
        # Also basename fallback for rel without prefix
        elif '/' not in rel:
            # Check any file with same basename tagged
            for k, v in tags_map.items():
                if os.path.basename(k) == rel and any(str(t).lower() == RAG_INDEX_TAG for t in (v.get('tags') or [])):
                    out.append(rel)
                    break
    # Deduplicate preserving order
    seen = set()
    uniq = []
    for r in out:
        if r not in seen:
            seen.add(r)
            uniq.append(r)
    return uniq


def should_inject_local_rag(task, project_path):
    """Return True only when rag-index opt-in is explicit (Block D1 contract).

    Injection fires when:
      * any context_refs entry has the ``rag-index`` tag, OR
      * task description contains ``rag:`` keyword (case-insensitive).
    Never auto-injected otherwise.
    """
    if _has_rag_keyword(task):
        return True
    if _task_rag_index_refs(task, project_path):
        return True
    return False


def retrieve_local_rag(task, project_path, project=None):
    """Retrieve host-local rag-index chunks for an opted-in task.

    Returns ``(block_text, provenance)``. Empty block when not opted in or on
    failure. ``provenance`` is ``{via: 'local-rag', rels, query, n_hits}``.
    """
    if not should_inject_local_rag(task, project_path):
        return '', {}
    # EU-only: local RAG is host-local under /opt/RAG, so allowed even for
    # eu_only projects — never calls a US model.
    desc = (task or {}).get('description') or ''
    # Extract rag: query if present
    query = ''
    low = desc.lower()
    if 'rag:' in low:
        idx = low.index('rag:')
        query = desc[idx + 4:].strip().split('\n')[0].strip()
        # Also include next sentence if short
        if len(query) < 5 and '\n' in desc[idx:]:
            query = desc[idx + 4: idx + 4 + 300].strip()
        query = query[:_RAG_QUERY_CHAR_CAP] if query else ''
    if not query:
        # Fallback to focused legal query helper
        try:
            query = _extract_legal_query(desc, (task or {}).get('title') or '')
        except Exception:
            query = (task or {}).get('title') or ''
    query = (query or '').strip()[:_RAG_QUERY_CHAR_CAP]

    rag_refs = _task_rag_index_refs(task, project_path)
    # If no tagged refs but rag: keyword is present, consider all rag-index files in project
    if not rag_refs and _has_rag_keyword(task):
        try:
            import agent_db as _db
            tags_map = _db.get_file_tags(project_path) or {}
            rag_refs = [k for k, v in tags_map.items() if any(str(t).lower() == RAG_INDEX_TAG for t in (v.get('tags') or []))]
            # Exclude hold-out files
            rag_refs = [r for r in rag_refs if not is_holdout_file(r)]
        except Exception:
            rag_refs = []

    # Exclude hold-out refs at retrieval time as well
    rag_refs = [r for r in rag_refs if not is_holdout_file(r)]

    citations = []
    for rel in rag_refs[:5]:  # cap files to avoid huge prompt
        idx_path = _rag_local_index_path(project_path, rel)
        if not os.path.isfile(idx_path):
            # Fallback: read the source file directly and chunk on the fly
            full = os.path.join(project_path, rel)
            if not os.path.isfile(full) and '/' not in rel:
                for folder in ('Working Documents', 'Working Docs', 'My Docs', 'working-docs', 'docs'):
                    cand = os.path.join(project_path, folder, rel)
                    if os.path.isfile(cand):
                        full = cand
                        break
            if not os.path.isfile(full):
                continue
            text = _rag_extract_text(full)
            if not text:
                continue
            chunks = _rag_split_text(text, _RAG_CHUNK_CHARS)
            # Simple keyword scoring if query present
            scored = []
            if query:
                q_tokens = set(re.findall(r'[a-ząćęłńóśźż0-9]+', query.lower()))
                for i, ch in enumerate(chunks):
                    toks = set(re.findall(r'[a-ząćęłńóśźż0-9]+', ch.lower()))
                    overlap = len(q_tokens & toks) if q_tokens else 0
                    scored.append((overlap, i, ch))
                scored.sort(key=lambda x: (-x[0], x[1]))
                # Take top 2 chunks per file
                for _, _, ch in scored[:2]:
                    citations.append({'rel': rel, 'extract': ch[:400], 'chunk_id': f'{rel}::chunk'})
            else:
                for ch in chunks[:2]:
                    citations.append({'rel': rel, 'extract': ch[:400], 'chunk_id': f'{rel}::chunk'})
            continue
        try:
            with open(idx_path, 'r', encoding='utf-8') as fh:
                data = json.load(fh)
            chunks = data.get('chunks') or []
            if query:
                q_tokens = set(re.findall(r'[a-ząćęłńóśźż0-9]+', query.lower()))
                scored = []
                for c in chunks:
                    txt = c.get('text') or ''
                    toks = set(re.findall(r'[a-ząćęłńóśźż0-9]+', txt.lower()))
                    overlap = len(q_tokens & toks) if q_tokens else 0
                    scored.append((overlap, txt))
                scored.sort(key=lambda x: -x[0])
                for _, txt in scored[:2]:
                    citations.append({'rel': rel, 'extract': txt[:400], 'chunk_id': f'{rel}::chunk'})
            else:
                for c in chunks[:2]:
                    citations.append({'rel': rel, 'extract': (c.get('text') or '')[:400], 'chunk_id': c.get('chunk_id')})
        except Exception as e:
            _log.warning('[rag-index] retrieve local read failed for %s: %s', rel, e)
            continue

    if not citations:
        return '', {'via': 'local-rag', 'rels': rag_refs, 'query': query, 'n_hits': 0}

    # Build inline block (host-local, no SCW)
    lines = [f'## Retrieved local context (RAG rag-index: {len(citations)} chunk(s))']
    lines.append(
        'This is the host-local rag-index library (/opt/RAG, on-host, no SCW). '
        'Answer ONLY from these excerpts when they are relevant; cite the file rel + verbatim extract.'
    )
    if query:
        lines.append(f'Query: {query}')
    for c in citations[:6]:
        rel = c.get('rel') or ''
        ext = c.get('extract') or ''
        lines.append(f'[{rel}] {ext}')
    block = '\n\n'.join(lines)[:_RAG_INLINE_CHAR_CAP]

    provenance = {
        'via': 'local-rag',
        'rels': rag_refs,
        'query': query,
        'n_hits': len(citations),
        'citation_ids': [c.get('chunk_id') for c in citations],
        'host_local': True,
        'rag_root': _RAG_LOCAL_ROOT,
    }
    return block, provenance
